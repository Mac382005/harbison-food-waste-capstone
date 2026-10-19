from datetime import date, timedelta

import pandas as pd
import pytest

import src.db as db
from src.intake import ensure_standard_orders, orders_for, residents
from src.synthetic_intake import generate_synthetic_resident_data
from src.resident_import import apply, normalise, preview

ROWS = [
    {"Resident ID": "T1001", "Name": "Ada Pemberton", "Room": "120", "Dining area": "Area A", "Menu type": "Regular"},
    {"Resident ID": "t1002", "Name": "Clive Hartnell", "Room": "121", "Dining area": "area b", "Menu type": ""},
]


def _plan(rows, source="fictional_demo"):
    return preview(pd.DataFrame(rows), source)


def test_new_residents_are_imported_and_listed(temp_db):
    plan = _plan(ROWS)
    assert plan.Action.tolist() == ["New", "New"]
    assert plan["Resident ID"].tolist() == ["T1001", "T1002"]  # IDs upper-cased
    assert apply(plan, "fictional_demo", None)["New"] == 2
    people = {p["resident_code"]: p for p in residents("fictional_demo")}
    assert people["T1002"]["display_name"] == "Clive Hartnell" and people["T1002"]["service_area"] == "Area B"
    assert people["T1002"]["menu_variant"] == "Regular"  # blank menu type defaults to Regular


def test_reupload_updates_room_and_discharges(temp_db):
    apply(_plan(ROWS), "fictional_demo", None)
    moved = [dict(ROWS[0], Room="130"), dict(ROWS[1], Status="Discharged")]
    plan = _plan(moved)
    assert plan.Action.tolist() == ["Update", "Discharge"]
    assert "room 120 → 130" in plan.Changes.iloc[0]
    apply(plan, "fictional_demo", None)
    active = {p["resident_code"]: p["room"] for p in residents("fictional_demo")}
    assert active == {"T1001": "130"}  # discharged resident hidden from the nurse list
    assert _plan(moved).Action.tolist() == ["No change", "No change"]
    assert _plan([dict(ROWS[0], Room="130"), dict(ROWS[1], Status="Active")]).Action.tolist() == ["No change", "Re-admit"]
    # a file without T1001 takes T1001 off the nurse list
    assert _plan([dict(ROWS[1], Status="Active")]).Action.tolist() == ["Re-admit", "Remove (not in file)"]


def test_problem_rows_are_flagged_not_saved(temp_db):
    plan = _plan([
        dict(ROWS[0]), dict(ROWS[0]),  # duplicate ID
        {"Resident ID": "T2", "Name": "", "Room": "", "Dining area": "Nowhere", "Menu type": "Liquid"},
        {"Resident ID": "DEMO-01-001", "Name": "X", "Room": "1", "Dining area": "Area A"},
        {"Resident ID": "T3003", "Name": "Y", "Room": "2", "Dining area": "Area A", "Status": "maybe"},
    ])
    assert (plan.Action == "Error").all()
    problems = " | ".join(plan.Problem)
    for text in ["more than once", "3-25 letters", "Name missing", "Room missing", "'Nowhere' not found",
                 "Diet 'Liquid' not recognised", "reserved", "Status 'maybe'"]:
        assert text in problems
    assert apply(plan, "fictional_demo", None)["New"] == 0
    assert residents("fictional_demo") == []


def test_missing_columns_and_alternative_headings(temp_db):
    with pytest.raises(ValueError, match="Missing columns: room"):
        preview(pd.DataFrame([{"Resident ID": "T1", "Name": "A", "Dining area": "Area A"}]), "fictional_demo")
    renamed, missing = normalise(pd.DataFrame([{"URN": "T9", "Resident Name": "B", "Bed": "7", "Wing": "Area A"}]))
    assert missing == [] and set(renamed.columns) == {"resident_id", "name", "room", "dining_area"}


def test_real_upload_needs_real_records_and_ids_cannot_switch_source(temp_db):
    apply(_plan(ROWS), "fictional_demo", None)
    assert "other data source" in _plan(ROWS[:1], source="real").Problem.iloc[0]


def test_uploaded_residents_get_standard_orders_and_survive_demo_regeneration(temp_db):
    apply(_plan(ROWS), "fictional_demo", None)
    day = date.today() - timedelta(days=1)
    people = residents("fictional_demo")
    assert ensure_standard_orders(people, day, "Lunch") == 2
    assert ensure_standard_orders(people, day, "Lunch") == 0  # only once
    lunch = {o["item_label"] for o in orders_for("fictional_demo", day, "Lunch")[people[0]["id"]]}
    assert lunch == {"Roast Chicken", "Mashed Potato", "Steamed Broccoli", "Apple Crumble"}
    ensure_standard_orders(people, day, "Morning Tea")
    tea = {o["item_label"] for o in orders_for("fictional_demo", day, "Morning Tea")[people[0]["id"]]}
    assert tea == {"Tea"}
    generate_synthetic_resident_data(day - timedelta(days=2), day, residents_per_area=2)
    codes = {p["resident_code"] for p in residents("fictional_demo")}
    assert {"T1001", "T1002"} <= codes and len(codes) == 6
    generate_synthetic_resident_data(day - timedelta(days=2), day, residents_per_area=1)
    assert {p["resident_code"] for p in residents("fictional_demo")} >= {"T1001", "T1002"}
    assert db.fetchone("SELECT COUNT(*) n FROM intake_residents WHERE resident_code LIKE 'DEMO-%'")["n"] == 2


def test_texture_diets_are_recognised_and_labelled(temp_db):
    from src.intake import menu_components_for_meal
    from src.textures import parse_texture, texture_label
    assert [parse_texture(x) for x in ["Minced and moist", "MM5", "Level 4", "Pureed", "Soft & Bite-Sized", ""]] ==         ["Minced", "Minced", "Puree", "Puree", "Soft Bite Sized", "Regular"]
    assert parse_texture("Liquid") is None
    assert texture_label("Minced", with_level=True) == "Minced & Moist (Level 5)"

    plan = _plan([dict(ROWS[0], **{"Menu type": "Minced and moist"}), dict(ROWS[1], **{"Menu type": "PU4"})])
    assert plan.Diet.tolist() == ["Minced & Moist", "Pureed"]
    apply(plan, "fictional_demo", None)
    assert {p["resident_code"]: p["menu_variant"] for p in residents("fictional_demo")} ==         {"T1001": "Minced", "T1002": "Puree"}
    day = date.today() - timedelta(days=1)
    labels = [i["item_label"] for i in menu_components_for_meal(day, "Lunch", "Minced")]
    assert "Roast Chicken (Minced & Moist)" in labels  # demo menu has no minced version, so it is labelled
    assert "Roast Chicken" in [i["item_label"] for i in menu_components_for_meal(day, "Lunch", "Regular")]
    assert "diet Minced & Moist → Pureed" in _plan([dict(ROWS[0], **{"Menu type": "Pureed"})]).Changes.iloc[0]


def test_generated_residents_cover_all_diets_without_melting_desserts(temp_db):
    day = date.today() - timedelta(days=1)
    generate_synthetic_resident_data(day - timedelta(days=6), day, residents_per_area=20)
    diets = {r["menu_variant"] for r in db.fetchall("SELECT menu_variant FROM intake_residents")}
    assert diets == {"Regular", "Soft Bite Sized", "Minced", "Puree"}
    bad = db.fetchone("""SELECT COUNT(*) n FROM intake_meal_orders o JOIN intake_residents r ON r.id=o.resident_id
                         WHERE r.menu_variant IN ('Minced','Puree') AND o.item_label IN ('Jelly','Ice cream')""")["n"]
    assert bad == 0
    minced = db.fetchone("""SELECT COUNT(*) n FROM intake_meal_orders o JOIN intake_residents r ON r.id=o.resident_id
                            WHERE r.menu_variant='Minced' AND o.item_label LIKE '%(Minced & Moist)'""")["n"]
    assert minced > 0


def test_harbison_soups_use_recipe_serving_size(temp_db, tmp_path, monkeypatch):
    import src.resident_choices as pc
    from src.intake import ComponentEntry, MealEntry, save_observation
    soups = tmp_path / "soups.csv"
    soups.write_text("soup_name,style,serves,serving_size_g,yield_loss_pct,source_page,source_document\n"
                     "Minestrone Soup,Chunky,100,240,10,3,Soups.pdf\nTomato Soup,Blended,100,240,20,11,Soups.pdf\n",
                     encoding="utf-8")
    monkeypatch.setattr(pc, "SOUPS_FILE", soups)
    ids = pc.ensure_soup_items()
    assert set(ids) == {"Minestrone Soup", "Tomato Soup"} and pc.ensure_soup_items() == ids  # idempotent
    apply(_plan(ROWS[:1]), "fictional_demo", None)
    rid = residents("fictional_demo")[0]["id"]
    day = date.today() - timedelta(days=1)
    save_observation(MealEntry(rid, day, "Lunch", "served", [
        ComponentEntry("soup", "Minestrone Soup", "observed", 50, ids["Minestrone Soup"], served_weight_g=240)]), None)
    comp = db.fetchone("SELECT served_weight_g,intake_pct FROM intake_component_observations")
    assert comp == {"served_weight_g": 240.0, "intake_pct": 50}
    assert db.fetchone("SELECT COUNT(*) n FROM transactions WHERE intake_observation_id IS NOT NULL")["n"] == 0


def test_add_item_options_hide_melting_desserts_for_pureed(temp_db):
    from views.plate_waste import _item_options
    day = date.today() - timedelta(days=1)
    regular = {o[1] for o in _item_options({"menu_variant": "Regular"}, day, "Lunch")}
    pureed = {o[1] for o in _item_options({"menu_variant": "Puree"}, day, "Lunch")}
    assert {"Roast Chicken", "Pavlova", "Biscuits", "Pineapple juice", "Vanilla ice cream"} <= regular
    assert not {"Chocolate cake", "Water"} & regular  # only items on the Harbison supplier list
    assert "Roast Chicken (Pureed)" in pureed and "Pavlova" in pureed
    assert not {"Vanilla ice cream", "Strawberry jelly"} & pureed


def test_nurse_can_correct_room_and_dining_area(temp_db):
    from src.intake import update_resident_location
    apply(_plan(ROWS), "fictional_demo", None)
    people = {p["resident_code"]: p for p in residents("fictional_demo")}
    ada, clive = people["T1001"], people["T1002"]
    assert update_resident_location(ada["id"], " 125 ", ada["service_area_id"], None) is None
    assert {p["resident_code"]: p["room"] for p in residents("fictional_demo")}["T1001"] == "125"
    warning = update_resident_location(ada["id"], "121", clive["service_area_id"], None)  # Clive's room in Area B
    assert "also assigned to Clive Hartnell" in warning
    moved = {p["resident_code"]: p for p in residents("fictional_demo")}["T1001"]
    assert moved["service_area"] == "Area B" and moved["room"] == "121"
    with pytest.raises(ValueError):
        update_resident_location(ada["id"], "  ", ada["service_area_id"], None)
    assert db.fetchone("SELECT COUNT(*) n FROM audit_log WHERE action='RESIDENT_MOVED'")["n"] == 2


def test_tea_time_offers_only_drinks_and_light_snacks(temp_db):
    from views.plate_waste import _item_options
    day = date.today() - timedelta(days=1)
    for meal in ("Morning Tea", "Afternoon Tea"):
        options = _item_options({"menu_variant": "Regular"}, day, meal)
        assert {o[0] for o in options} == {"Drinks", "Desserts", "Biscuits & snacks", "Other"}
        assert not {"Sausage roll", "Beef pie", "Chicken schnitzel", "Roast Chicken"} & {o[1] for o in options}
    for meal in ("Lunch", "Dinner"):
        groups = {o[0] for o in _item_options({"menu_variant": "Regular"}, day, meal)}
        assert {"Hot food", "Soups", "Today's menu"} <= groups and "Breakfast" not in groups
    breakfast = {o[0] for o in _item_options({"menu_variant": "Regular"}, day, "Breakfast")}
    assert "Breakfast" in breakfast and not {"Hot food", "Soups"} & breakfast


def test_upload_replaces_the_nurse_list_but_keeps_history(temp_db):
    from src.intake import ComponentEntry, MealEntry, save_observation
    day = date.today() - timedelta(days=1)
    generate_synthetic_resident_data(day - timedelta(days=2), day, residents_per_area=2)
    demo = residents("fictional_demo")
    assert len(demo) == 4
    history = db.fetchone("SELECT COUNT(*) n FROM intake_meal_observations")["n"]
    plan = _plan(ROWS)
    assert plan.Action.value_counts().to_dict() == {"New": 2, "Remove (not in file)": 4}
    done = apply(plan, "fictional_demo", None)
    assert done["New"] == 2 and done["Remove (not in file)"] == 4
    assert {p["resident_code"] for p in residents("fictional_demo")} == {"T1001", "T1002"}  # only the upload
    assert db.fetchone("SELECT COUNT(*) n FROM intake_meal_observations")["n"] == history  # history kept
    assert len(residents("fictional_demo", active_only=False)) == 6


def test_reasons_and_notes_are_visible_to_dietitians(temp_db):
    from src.intake import ComponentEntry, MealEntry, save_observation
    from views.reasons_notes import _load
    apply(_plan(ROWS), "fictional_demo", None)
    rid = {p["resident_code"]: p["id"] for p in residents("fictional_demo")}["T1001"]
    day = date.today() - timedelta(days=1)
    save_observation(MealEntry(rid, day, "Lunch", "served", [
        ComponentEntry("soup", "Pumpkin Soup", "observed", 25, None, "Poor appetite", served_weight_g=250),
        ComponentEntry("protein", "Roast Chicken", "observed", 100, served_weight_g=180),
        ComponentEntry("dessert", "Apple Crumble", "observed", 50, served_weight_g=120),
    ], staff_observation="Said the soup was too salty"), None)
    items, notes = _load(day, day, None, "fictional_demo")
    assert items[["Item", "Eaten", "Reason"]].values.tolist() == [
        ["Pumpkin Soup", "25%", "Poor appetite"], ["Apple Crumble", "50%", ""]]  # 100% eaten is left out
    assert items["Resident"].tolist() == ["Ada Pemberton", "Ada Pemberton"]
    assert notes[["Resident", "Meal", "Note"]].values.tolist() == [["Ada Pemberton", "Lunch", "Said the soup was too salty"]]


def test_excel_report_is_formatted_for_dietitians():
    import io
    import openpyxl
    from src.intake_excel import build_report
    items = pd.DataFrame([
        {"Date": "2026-10-16", "Meal": "Lunch", "Resident": "Ada Pemberton", "Room": "120", "Diet": "Pureed",
         "Item": "Pumpkin Soup", "Eaten": "25%", "Reason": "Poor appetite", "Recorded by": "Nurse"},
        {"Date": "2026-10-16", "Meal": "Dinner", "Resident": "Ada Pemberton", "Room": "120", "Diet": "Pureed",
         "Item": "Custard", "Eaten": "50%", "Reason": "", "Recorded by": "Nurse"}])
    notes = pd.DataFrame([{"Date": "2026-10-16", "Meal": "Lunch", "Resident": "Ada Pemberton", "Room": "120",
                           "Outcome": "Ate", "Note": "Too salty", "Written by": "Nurse"}])
    wb = openpyxl.load_workbook(io.BytesIO(build_report(items, notes, {"period": "16 Oct", "generated_by": "Dietitian"})))
    assert wb.sheetnames == ["Summary", "Low-intake items", "Nurse notes"]
    ws = wb["Low-intake items"]
    assert [c.value for c in ws[1]] == ["Date", "Meal", "Resident", "Room", "Diet", "Item", "% eaten", "Reason", "Recorded by"]
    assert ws["G2"].value == 0.25 and ws["G2"].number_format == "0%" and ws["H3"].value == "Not given"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    summary = [c.value for row in wb["Summary"].iter_rows() for c in row if c.value is not None]
    assert "Poor appetite" in summary and "Ada Pemberton" in summary and "Dietitian" in summary


def test_upload_for_one_site_leaves_other_sites_alone(temp_db):
    with db.connection() as c:
        c.execute("INSERT INTO facilities(name) VALUES('Second Site')")
        c.execute("INSERT INTO service_areas(facility_id,name) VALUES(2,'Area A')")
    site1, site2 = 1, 2
    apply(preview(pd.DataFrame(ROWS), "fictional_demo", facility_id=site1), "fictional_demo", None)
    other = [{"Resident ID": "S2-001", "Name": "Maud Keller", "Room": "301", "Dining area": "Area A"}]
    apply(preview(pd.DataFrame(other), "fictional_demo", facility_id=site2), "fictional_demo", None)
    people = {p["resident_code"]: p for p in residents("fictional_demo")}
    assert set(people) == {"T1001", "T1002", "S2-001"}  # site 2 upload did not remove site 1 residents
    assert people["S2-001"]["facility"] == "Second Site"
    # a new site-1 list without T1002 removes only T1002
    plan = preview(pd.DataFrame(ROWS[:1]), "fictional_demo", facility_id=site1)
    assert sorted(plan.Action) == ["No change", "Remove (not in file)"]
    apply(plan, "fictional_demo", None)
    assert {p["resident_code"] for p in residents("fictional_demo")} == {"T1001", "S2-001"}
    # rows naming another facility are flagged
    wrong = preview(pd.DataFrame([dict(ROWS[0], Facility="Second Site")]), "fictional_demo", facility_id=site1)
    assert "not the selected site" in wrong.Problem.iloc[0]
    # 'Area B' exists only at site 1, so it is not found when site 2 is selected
    missing = preview(pd.DataFrame([dict(ROWS[1], **{"Resident ID": "S2-002"})]), "fictional_demo", facility_id=site2)
    assert "not found at Second Site" in missing.Problem.iloc[0]


def test_file_with_no_usable_rows_never_removes_anyone(temp_db):
    apply(_plan(ROWS), "fictional_demo", None)
    wrong_site_file = [{"Resident ID": "X9-001", "Name": "Someone", "Room": "1", "Dining area": "Nowhere"}]
    plan = _plan(wrong_site_file)
    assert plan.Action.tolist() == ["Error"]  # no "Remove (not in file)" rows at all
    apply(plan, "fictional_demo", None)
    assert {p["resident_code"] for p in residents("fictional_demo")} == {"T1001", "T1002"}
